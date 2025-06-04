import os
import shutil
from uuid import uuid4
from dotenv import load_dotenv

from docx import Document
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_openai.embeddings import OpenAIEmbeddings
from langchain_chroma import Chroma
from langchain.docstore.document import Document as LangDocument

import re

# Load môi trường
load_dotenv()

# Cấu hình
DATA_PATH = r"data"
CHROMA_PATH = r"chroma_db"

# Xoá dữ liệu cũ
shutil.rmtree(CHROMA_PATH, ignore_errors=True)
print("✅ Đã xoá dữ liệu cũ trong ChromaDB.")

# Embedding model
embeddings_model = OpenAIEmbeddings(model="text-embedding-3-large")

# Khởi tạo vector store
vector_store = Chroma(
    collection_name="example_collection",
    embedding_function=embeddings_model,
    persist_directory=CHROMA_PATH,
)

# ---------------------------------------------------
# Hàm xử lý .docx thành văn bản markdown có heading, code, đoạn thường
def extract_structured_text(docx_path):
    doc = Document(docx_path)
    result = []
    code_block = []  # Danh sách để gom các dòng code
    in_code_block = False  # Biến kiểm tra có đang ở trong đoạn code hay không

    for para in doc.paragraphs:
        style = para.style.name.lower()
        text = para.text.strip()

        if not text:
            continue

        # Nếu đoạn này là code
        if "code" in style:
            if not in_code_block:
                in_code_block = True
                code_block = []
            code_block.append(text)
        else:
            # Nếu đang trong code block mà gặp đoạn thường, thì kết thúc code block
            if in_code_block:
                result.append("```c\n" + "\n".join(code_block) + "\n```")
                in_code_block = False

            if "heading" in style:
                level = int(re.search(r'\d+', style).group()) if re.search(r'\d+', style) else 2
                result.append(f"{'#' * level} {text}")

            else:
                result.append(text)  # Paragraph thường

    # Nếu file kết thúc trong một đoạn code block
    if in_code_block:
        result.append("```c\n" + "\n".join(code_block) + "\n```")

    return "\n\n".join(result)

# ---------------------------------------------------
# Hàm tách theo heading "##" → mỗi section là 1 chunk
def split_by_heading(document: LangDocument):
    sections = re.split(r'^(#{2,6})\s+', document.page_content, flags=re.MULTILINE)
    chunks = []
    for section in sections:
        if not section.strip():
            continue
        lines = section.strip().split('\n', 1)
        heading = lines[0].strip()
        content = lines[1].strip() if len(lines) > 1 else ""

        full_chunk = f"## {heading}\n\n{content}"

        chunks.append(LangDocument(
            page_content=full_chunk,
            metadata={
                "source": document.metadata["source"],
                "heading": heading
            }
        ))
    return chunks

# ---------------------------------------------------
# Tiền xử lý tất cả .docx
documents = []
for filename in os.listdir(DATA_PATH):
    if filename.endswith(".docx"):
        filepath = os.path.join(DATA_PATH, filename)
        structured_text = extract_structured_text(filepath)
        doc = LangDocument(page_content=structured_text, metadata={"source": filename})
        documents.append(doc)

print(f"📄 Đã load {len(documents)} tài liệu Word.")

# Tách theo heading
all_chunks = []
for doc in documents:
    all_chunks.extend(split_by_heading(doc))

print(f"🔹 Đã tách thành {len(all_chunks)} đoạn theo heading.")

# ---------------------------------------------------
# Với các đoạn dài quá → tách tiếp bằng splitter nhỏ hơn
final_chunks = []
splitter = RecursiveCharacterTextSplitter(
    chunk_size=1500,
    chunk_overlap=200
)
chunkss = []
for chunk in all_chunks:
    if len(chunk.page_content) > 1500:
        chunkss.append(chunk)
        sub_chunks = splitter.split_documents([chunk])
        for sub_chunk in sub_chunks:
            sub_chunk.metadata.update(chunk.metadata)  # Giữ metadata gốc
            final_chunks.append(sub_chunk)
    else:
        final_chunks.append(chunk)
for chunk in chunkss:
    if len(chunk.page_content) == 9175:
        print(chunk.page_content)
print(f"🔧 Tổng cộng {len(final_chunks)} chunks sau khi xử lý tách nhỏ.")

# ---------------------------------------------------
# Thêm vào vector DB
uuids = [str(uuid4()) for _ in range(len(final_chunks))]
vector_store.add_documents(documents=final_chunks, ids=uuids)

print("✅ Đã thêm vào vector store và lưu ChromaDB!")
